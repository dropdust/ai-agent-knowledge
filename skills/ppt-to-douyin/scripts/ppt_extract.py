#!/usr/bin/env python3
"""Extract slide text from a .pptx into segments.json for the Douyin pipeline.

Each segment: {index, title, bullets[], narration, notes}
- title: slide title placeholder text, else the first text box
- bullets: remaining paragraph lines (order preserved)
- narration: title + bullets joined with Chinese punctuation -- the text TTS reads
"""
import argparse
import json

from pptx import Presentation


def _clean(s):
    return s.replace("\u3000", " ").strip()


def extract(pptx_path, include_notes=False):
    prs = Presentation(pptx_path)
    segments = []
    for i, slide in enumerate(prs.slides, 1):
        title = None
        try:
            if slide.shapes.title is not None:
                title = _clean(slide.shapes.title.text)
        except Exception:
            title = None

        bullets = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                t = _clean("".join(r.text for r in p.runs))
                if not t:
                    continue
                if title is None:
                    title = t
                elif t != title:
                    bullets.append(t)

        notes = ""
        if include_notes:
            try:
                if slide.has_notes_slide:
                    notes = _clean(slide.notes_slide.notes_text_frame.text)
            except Exception:
                notes = ""

        narration = (title + "。" if title else "") + "，".join(bullets) + "。"

        segments.append({
            "index": i,
            "title": title or "",
            "bullets": bullets,
            "narration": narration,
            "notes": notes,
        })
    return segments


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("input", help="path to .pptx")
    ap.add_argument("-o", "--output", default="segments.json")
    ap.add_argument("--include-notes", action="store_true", help="also extract speaker notes")
    args = ap.parse_args()

    segments = extract(args.input, args.include_notes)
    with open(args.output, "w", encoding="utf-8") as f:
        json.dump({"segments": segments}, f, ensure_ascii=False, indent=2)
    print(f"Wrote {len(segments)} segments -> {args.output}")
    for s in segments:
        print(f"  [{s['index']:02d}] {s['title']}")


if __name__ == "__main__":
    main()
